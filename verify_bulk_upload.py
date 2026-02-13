
import os
import django
import sys

# Setup Django environment
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'school_sms.settings')
django.setup()

from django.core.files.uploadedfile import SimpleUploadedFile
from dashboards.bulk_operations import BulkImporter
from dashboards.models import QuestionBank
from django.contrib.auth import get_user_model
import openpyxl
import docx
import io

User = get_user_model()

def verify_excel_import():
    print("Verifying Excel Import...")
    # Create Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = [
        'text', 'type', 'marks', 'difficulty', 
        'choice_1', 'correct_1', 
        'choice_2', 'correct_2'
    ]
    ws.append(headers)
    ws.append([
        "Excel Test Question", "objective", "5", "hard",
        "Option A", "False",
        "Option B", "True"
    ])
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    file = SimpleUploadedFile("test_questions.xlsx", buffer.getvalue())
    
    # Run import
    admin = User.objects.filter(is_superuser=True).first()
    job = BulkImporter.import_questions(file, admin)
    
    if job.status == 'completed':
        print("Excel Import Job Completed")
        q = QuestionBank.objects.filter(text="Excel Test Question").first()
        if q:
            print("Excel Question Found: OK")
        else:
            print("Excel Question Found: FAIL")
    else:
        print(f"Excel Import Failed: {job.error_log}")

def verify_word_import():
    print("Verifying Word Import...")
    # Create Word file
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=8)
    
    headers = [
        'text', 'type', 'marks', 'difficulty', 
        'choice_1', 'correct_1', 
        'choice_2', 'correct_2'
    ]
    
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        
    row_cells = table.add_row().cells
    data = [
        "Word Test Question", "objective", "5", "medium",
        "Option X", "True",
        "Option Y", "False"
    ]
    for i, d in enumerate(data):
        row_cells[i].text = str(d)
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    file = SimpleUploadedFile("test_questions.docx", buffer.getvalue())
    
    # Run import
    admin = User.objects.filter(is_superuser=True).first()
    job = BulkImporter.import_questions(file, admin)
    
    if job.status == 'completed':
        print("Word Import Job Completed")
        q = QuestionBank.objects.filter(text="Word Test Question").first()
        if q:
            print("Word Question Found: OK")
        else:
            print("Word Question Found: FAIL")
    else:
        print(f"Word Import Failed: {job.error_log}")

if __name__ == "__main__":
    try:
        verify_excel_import()
        verify_word_import()
    except Exception as e:
        print(f"Verification Error: {e}")
